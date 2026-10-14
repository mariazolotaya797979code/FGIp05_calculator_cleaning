"""Локальный калькулятор стоимости клининга. Работает без интернета."""

from __future__ import annotations

import ctypes
import math
import sys
import tkinter as tk
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tkinter import font as tkfont
from tkinter import messagebox

import customtkinter as ctk
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, ImageDraw, ImageFilter

NAVY = "#0f172a"
GLASS = "#173044"
GLASS_SOFT = "#1c3d4f"
MINT = "#a7f3d0"
MINT_DIM = "#6ee7b7"
TEAL = "#0d9488"
TEAL_DEEP = "#0f766e"
AQUA = "#2dd4bf"
TEXT = "#f8fafc"
MUTED = "#d7efe8"
FIELD = "#0c1c2e"
SHIMMER = (
    (167, 243, 208),
    (110, 231, 183),
    (94, 234, 212),
    (45, 212, 191),
    (20, 184, 166),
    (13, 148, 136),
    (15, 118, 110),
)

WEAR = 0.01
PERCENT_KEYS = (
    "Налог_Процент",
    "Маржа_Процент",
    "Наценка_выходные",
    "Наценка_праздники",
)


class DataError(Exception):
    pass


@dataclass
class Material:
    name: str
    unit: str
    price: float
    norm: float


@dataclass
class Service:
    type_name: str
    hourly_rate: float
    productivity: float
    complexity: float
    min_hours: float


@dataclass
class Constant:
    key: str
    value: float
    unit: str
    description: str


@dataclass
class Catalog:
    materials: list[Material]
    services: list[Service]
    constants: list[Constant]
    by_key: dict[str, Constant]
    loaded_at: datetime


@dataclass
class Result:
    materials: float
    labor: float
    extras: float
    base: float
    transport: float
    amortization: float
    taxes: float
    overhead: float
    margin: float
    weekend_fee: float
    holiday_fee: float
    markup: float
    total: float
    hours: float


def app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def data_file() -> Path:
    return app_dir() / "calculator_data.xlsx"


def money(value: float) -> str:
    sign = "−" if value < 0 else ""
    text = f"{abs(value):,.2f}".replace(",", " ").replace(".", ",")
    return f"{sign}{text} ₽"


def fmt_num(value: float, digits: int = 2) -> str:
    return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")


def round2(value: float) -> float:
    return round(value + 1e-9, 2)


def norm_header(value: object) -> str:
    return " ".join(str(value).replace("\n", " ").split()).casefold()


def norm_unit(value: str) -> str:
    return value.strip().casefold().rstrip(".")


def parse_number(value: object, field: str) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise DataError(f"Пустое значение в поле «{field}»")
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise DataError(f"Нечисловое значение в поле «{field}»")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
    if text.endswith("%"):
        text = text[:-1]
        try:
            return float(text) / 100
        except ValueError as exc:
            raise DataError(f"Нечисловое значение «{value}» в поле «{field}»") from exc
    try:
        return float(text)
    except ValueError as exc:
        raise DataError(f"Нечисловое значение «{value}» в поле «{field}»") from exc


def as_fraction(value: object, field: str) -> float:
    number = parse_number(value, field)
    if number > 1:
        number /= 100
    if number < 0:
        raise DataError(f"Процент «{field}» не может быть отрицательным")
    return number


def find_sheet(workbook, expected: str):
    for name in workbook.sheetnames:
        if name.strip().casefold() == expected.casefold():
            return workbook[name]
    names = ", ".join(workbook.sheetnames)
    raise DataError(f"В файле нет листа «{expected}». Сейчас в книге: {names}")


def header_map(sheet, expected: list[str]) -> tuple[int, dict[str, int]]:
    wanted = [norm_header(name) for name in expected]
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
        found: dict[str, int] = {}
        for index, cell in enumerate(row):
            if cell is None or str(cell).strip() == "":
                continue
            found[norm_header(cell)] = index
        if all(name in found for name in wanted):
            return row_index, found
    joined = ", ".join(expected)
    raise DataError(f"На листе «{sheet.title.strip()}» нет столбцов: {joined}")


def cell(row: tuple, headers: dict[str, int], name: str):
    index = headers[norm_header(name)]
    if index >= len(row):
        return None
    return row[index]


def _friendly_io_error(exc: Exception) -> DataError:
    text = str(exc).lower()
    winerror = getattr(exc, "winerror", None)
    if isinstance(exc, PermissionError) or winerror in {32, 33} or "permission" in text or "занят" in text:
        return DataError(
            "Файл calculator_data.xlsx открыт в Excel или занят другой программой. "
            "Закройте его и нажмите «Обновить данные»."
        )
    if isinstance(exc, FileNotFoundError):
        return DataError(
            "Не найден файл calculator_data.xlsx. Положите его в ту же папку, что и программа."
        )
    if isinstance(exc, (InvalidFileException, OSError)):
        return DataError(f"Не удалось прочитать calculator_data.xlsx. {exc}")
    return DataError(f"Не удалось прочитать calculator_data.xlsx. {exc}")


def load_catalog(path: Path | None = None) -> Catalog:
    source = path or data_file()
    if not source.exists():
        raise DataError(
            "Не найден файл calculator_data.xlsx. Положите его в ту же папку, что и программа."
        )
    try:
        workbook = load_workbook(source, data_only=True, read_only=True)
    except Exception as exc:
        raise _friendly_io_error(exc) from exc
    try:
        materials = _read_materials(find_sheet(workbook, "Materials"))
        services = _read_services(find_sheet(workbook, "Services"))
        constants = _read_constants(find_sheet(workbook, "Лист3"))
    finally:
        workbook.close()
    if not materials:
        raise DataError("На листе Materials нет ни одной строки с материалом")
    if not services:
        raise DataError("На листе Services нет ни одного типа услуги")
    by_key = {item.key: item for item in constants}
    missing = [key for key in ("Транспорт_Центр", "Транспорт_Область", *PERCENT_KEYS) if key not in by_key]
    if missing:
        raise DataError("На листе «Лист3» нет ключей: " + ", ".join(missing))
    return Catalog(materials, services, constants, by_key, datetime.now())


def _read_materials(sheet) -> list[Material]:
    header_row, headers = header_map(sheet, ["Название", "Ед_изм", "Цена_за_ед", "Норма_на_квм"])
    materials: list[Material] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if row is None or all(item is None or str(item).strip() == "" for item in row):
            continue
        name = cell(row, headers, "Название")
        if name is None or str(name).strip() == "":
            continue
        unit = cell(row, headers, "Ед_изм")
        materials.append(
            Material(
                name=str(name).strip(),
                unit="" if unit is None else str(unit).strip(),
                price=parse_number(cell(row, headers, "Цена_за_ед"), f"Цена_за_ед, строка {row_number}"),
                norm=parse_number(cell(row, headers, "Норма_на_квм"), f"Норма_на_квм, строка {row_number}"),
            )
        )
    return materials


def _read_services(sheet) -> list[Service]:
    header_row, headers = header_map(
        sheet,
        ["Тип_услуги", "Ставка в час", "Производительность кв/м в час", "Коэффициент сложности", "Минимум часов"],
    )
    services: list[Service] = []
    seen: dict[str, int] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if row is None or all(item is None or str(item).strip() == "" for item in row):
            continue
        raw_name = cell(row, headers, "Тип_услуги")
        if raw_name is None or str(raw_name).strip() == "":
            continue
        name = str(raw_name).strip()
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = f"{name} ({seen[name]})"
        productivity = parse_number(
            cell(row, headers, "Производительность кв/м в час"),
            f"Производительность, строка {row_number}",
        )
        min_hours = parse_number(cell(row, headers, "Минимум часов"), f"Минимум часов, строка {row_number}")
        if productivity <= 0:
            raise DataError(f"Производительность в строке {row_number} должна быть больше нуля")
        if min_hours < 0:
            raise DataError(f"Минимум часов в строке {row_number} не может быть отрицательным")
        services.append(
            Service(
                type_name=name,
                hourly_rate=parse_number(cell(row, headers, "Ставка в час"), f"Ставка в час, строка {row_number}"),
                productivity=productivity,
                complexity=parse_number(
                    cell(row, headers, "Коэффициент сложности"),
                    f"Коэффициент сложности, строка {row_number}",
                ),
                min_hours=min_hours,
            )
        )
    return services


def _read_constants(sheet) -> list[Constant]:
    header_row, headers = header_map(sheet, ["Ключ", "Значение (Цена за ед.)", "Ед. измерения", "Описание"])
    constants: list[Constant] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if row is None or all(item is None or str(item).strip() == "" for item in row):
            continue
        key = cell(row, headers, "Ключ")
        if key is None or str(key).strip() == "":
            continue
        key_text = str(key).strip()
        raw_value = cell(row, headers, "Значение (Цена за ед.)")
        value = as_fraction(raw_value, key_text) if key_text in PERCENT_KEYS else parse_number(raw_value, key_text)
        unit = cell(row, headers, "Ед. измерения")
        description = cell(row, headers, "Описание")
        constants.append(
            Constant(
                key=key_text,
                value=value,
                unit="" if unit is None else str(unit).strip(),
                description=key_text if description is None or str(description).strip() == "" else str(description).strip(),
            )
        )
    return constants


OPTIONAL_EXTRA_KEYS = ("Срочный_выезд", "Аренда_инвентаря")


def extra_services(catalog: Catalog) -> list[Constant]:
    return [
        item
        for item in catalog.constants
        if item.key.startswith("Цена_") or item.key in OPTIONAL_EXTRA_KEYS
    ]


def calculate(
    catalog: Catalog,
    area: float,
    service_name: str,
    quantities: dict[str, float],
    weekend: bool,
    holiday: bool,
    location: str,
) -> Result:
    if area <= 0:
        raise DataError("Площадь должна быть больше нуля")
    service = next((item for item in catalog.services if item.type_name == service_name), None)
    if service is None:
        raise DataError("Выберите тип уборки")
    if location not in {"Центр", "Область"}:
        raise DataError("Выберите локацию: Центр или Область")

    per_sqm = sum(item.norm * item.price for item in catalog.materials)
    materials = round2(area * per_sqm)
    hours = max(service.min_hours, area / service.productivity)
    labor = round2(hours * service.hourly_rate * service.complexity)

    extras_sum = 0.0
    known = {item.key: item for item in extra_services(catalog)}
    for key, quantity in quantities.items():
        if quantity < 0:
            raise DataError("Количество дополнительной услуги не может быть отрицательным")
        if quantity == 0:
            continue
        item = known.get(key)
        if item is None:
            raise DataError(f"В файле больше нет услуги «{key}». Нажмите «Обновить данные».")
        extras_sum += round2(item.value * quantity)
    extras_sum = round2(extras_sum)

    base = round2(materials + labor + extras_sum)
    transport_key = "Транспорт_Центр" if location == "Центр" else "Транспорт_Область"
    transport = round2(catalog.by_key[transport_key].value)
    inventory = sum(item.price for item in catalog.materials if norm_unit(item.unit) in {"шт", "час"})
    amortization = round2(inventory * WEAR)
    taxes = round2(base * catalog.by_key["Налог_Процент"].value)
    overhead = round2(transport + amortization + taxes)

    margin = round2(base * catalog.by_key["Маржа_Процент"].value)
    weekend_fee = round2(base * catalog.by_key["Наценка_выходные"].value) if weekend else 0.0
    holiday_fee = round2(base * catalog.by_key["Наценка_праздники"].value) if holiday else 0.0
    markup = round2(margin + weekend_fee + holiday_fee)
    total = round2(base + overhead + markup)
    return Result(
        materials=materials,
        labor=labor,
        extras=extras_sum,
        base=base,
        transport=transport,
        amortization=amortization,
        taxes=taxes,
        overhead=overhead,
        margin=margin,
        weekend_fee=weekend_fee,
        holiday_fee=holiday_fee,
        markup=markup,
        total=total,
        hours=hours,
    )


def render_icon(size: int = 512) -> Image.Image:
    mask = Image.new("L", (size, size), 0)
    ImageDraw.Draw(mask).rounded_rectangle((8, 8, size - 9, size - 9), radius=int(size * 0.24), fill=255)
    horizontal = Image.linear_gradient("L").resize((size, size))
    vertical = Image.linear_gradient("L").resize((size, size)).transpose(Image.Transpose.ROTATE_270)
    diagonal = Image.blend(horizontal, vertical, 0.45)
    mint = Image.new("RGBA", (size, size), (167, 243, 208, 255))
    teal = Image.new("RGBA", (size, size), (13, 148, 136, 255))
    icon = Image.composite(teal, mint, diagonal)
    gloss = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(gloss)
    draw.ellipse((size * 0.08, size * 0.02, size * 0.92, size * 0.62), fill=(255, 255, 255, 70))
    gloss = gloss.filter(ImageFilter.GaussianBlur(int(size * 0.04)))
    icon = Image.alpha_composite(icon, gloss)
    draw = ImageDraw.Draw(icon)
    cx, cy = size / 2, size * 0.56
    outer, inner = size * 0.24, size * 0.09
    points: list[tuple[float, float]] = []
    for index in range(8):
        angle = -math.pi / 2 + index * math.pi / 4
        radius = outer if index % 2 == 0 else inner
        points.append((cx + radius * math.cos(angle), cy + radius * math.sin(angle)))
    draw.polygon(points, fill=(255, 255, 255, 235))
    draw.ellipse((cx - size * 0.045, size * 0.16, cx + size * 0.045, size * 0.25), fill=(255, 255, 255, 230))
    icon.putalpha(mask)
    return icon


def ensure_brand_icons(folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    png_path = folder / "icon.png"
    ico_path = folder / "icon.ico"
    if png_path.exists() and ico_path.exists():
        return ico_path
    image = render_icon(512)
    image.save(png_path, "PNG")
    image.save(
        ico_path,
        format="ICO",
        sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)],
    )
    return ico_path


def shimmer_color(phase: float) -> str:
    spot = ((math.sin(phase) + 1) / 2) * (len(SHIMMER) - 1)
    left = int(spot)
    right = min(left + 1, len(SHIMMER) - 1)
    blend = spot - left
    rgb = tuple(int(SHIMMER[left][channel] * (1 - blend) + SHIMMER[right][channel] * blend) for channel in range(3))
    return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"


def render_background_layers(width: int, height: int) -> tuple[Image.Image, Image.Image]:
    width = max(width, 200)
    height = max(height, 200)
    top = Image.new("RGBA", (width, height), (15, 23, 42, 255))
    bottom = Image.new("RGBA", (width, height), (15, 61, 74, 255))
    mask = Image.linear_gradient("L").resize((height, width)).transpose(Image.Transpose.ROTATE_270)
    base = Image.composite(bottom, top, mask)
    sheet_w, sheet_h = int(width * 1.7), int(height * 1.7)
    glow = Image.new("RGBA", (sheet_w, sheet_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(glow)
    blooms = (
        (0.08, 0.02, 0.72, 0.55, (167, 243, 208, 78)),
        (0.42, 0.18, 1.15, 0.78, (110, 231, 183, 62)),
        (0.55, 0.00, 1.25, 0.48, (45, 212, 191, 70)),
        (-0.15, 0.35, 0.48, 1.05, (13, 148, 136, 66)),
        (0.25, 0.55, 0.95, 1.25, (20, 184, 166, 48)),
        (0.70, 0.48, 1.35, 1.15, (15, 118, 110, 58)),
    )
    for x0, y0, x1, y1, color in blooms:
        draw.ellipse((sheet_w * x0, sheet_h * y0, sheet_w * x1, sheet_h * y1), fill=color)
    return base, glow.filter(ImageFilter.GaussianBlur(78))


def compose_background(base: Image.Image, glow: Image.Image, phase: float) -> Image.Image:
    width, height = base.size
    glow_w, glow_h = glow.size
    span_x = max(glow_w - width, 1)
    span_y = max(glow_h - height, 1)
    origin_x = int((0.5 + 0.5 * math.sin(phase)) * span_x * 0.62)
    origin_y = int((0.5 + 0.5 * math.cos(phase * 0.73)) * span_y * 0.62)
    shifted = glow.crop((origin_x, origin_y, origin_x + width, origin_y + height))
    return Image.alpha_composite(base, shifted).convert("RGB")


class ExtraRow:
    def __init__(self, key: str, checked: tk.BooleanVar, quantity: tk.StringVar) -> None:
        self.key = key
        self.checked = checked
        self.quantity = quantity


class CalculatorApp(ctk.CTk):
    def __init__(self) -> None:
        super().__init__(fg_color=NAVY)
        ctk.set_appearance_mode("dark")
        self.title("Калькулятор клининга")
        self.minsize(1080, 720)
        self.catalog: Catalog | None = None
        self.extra_rows: list[ExtraRow] = []
        self._bg_size = (0, 0)
        self._resize_job: str | None = None
        self._base_layer: Image.Image | None = None
        self._glow_layer: Image.Image | None = None
        self._phase = 0.0
        self._cards: list[ctk.CTkFrame] = []
        self._family = "Segoe UI"

        self.area_var = tk.StringVar(value="50")
        self.service_var = tk.StringVar()
        self.location_var = tk.StringVar(value="Центр")
        self.weekend_var = tk.BooleanVar(value=False)
        self.holiday_var = tk.BooleanVar(value=False)
        self.service_hint = tk.StringVar(value="")
        self.location_hint = tk.StringVar(value="")
        self.status_var = tk.StringVar(value="Загрузка данных…")
        self.base_var = tk.StringVar(value="—")
        self.overhead_var = tk.StringVar(value="—")
        self.markup_var = tk.StringVar(value="—")
        self.total_var = tk.StringVar(value="—")
        self.base_detail = tk.StringVar(value="")
        self.overhead_detail = tk.StringVar(value="")
        self.markup_detail = tk.StringVar(value="")

        self._fit_to_screen()
        self._apply_icon()
        self._build()
        self.bind("<Configure>", self._on_resize)
        self.after(30, self._paint_background)
        self.after(140, self._shimmer)
        self.reload_data(announce=False)

    def _fit_to_screen(self) -> None:
        self.update_idletasks()
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        width = min(1240, max(980, int(screen_w * 0.9)))
        height = min(860, max(680, int(screen_h * 0.86)))
        width = min(width, screen_w - 24)
        height = min(height, screen_h - 72)
        self.minsize(min(980, width), min(640, height))
        x = max((screen_w - width) // 2, 0)
        y = max((screen_h - height) // 2, 0)
        self.geometry(f"{width}x{height}+{x}+{y}")

    def _font(self, size: int, weight: str = "normal") -> ctk.CTkFont:
        return ctk.CTkFont(family=self._family, size=size, weight=weight)

    def _apply_icon(self) -> None:
        try:
            families = set(tkfont.families(self))
            if "Bahnschrift" in families:
                self._family = "Bahnschrift"
            ico = ensure_brand_icons(app_dir())
            self.iconbitmap(str(ico))
            png = Image.open(app_dir() / "icon.png")
            self.logo_image = ctk.CTkImage(light_image=png, dark_image=png, size=(42, 42))
        except Exception:
            self.logo_image = None

    def _on_resize(self, event) -> None:
        if event.widget is not self:
            return
        if self._resize_job is not None:
            self.after_cancel(self._resize_job)
        self._resize_job = self.after(90, self._paint_background)

    def _paint_background(self) -> None:
        self._resize_job = None
        width, height = self.winfo_width(), self.winfo_height()
        if width < 100 or height < 100 or (width, height) == self._bg_size:
            return
        self._bg_size = (width, height)
        self._base_layer, self._glow_layer = render_background_layers(width, height)
        self._show_background()

    def _show_background(self) -> None:
        if self._base_layer is None or self._glow_layer is None:
            return
        image = compose_background(self._base_layer, self._glow_layer, self._phase)
        width, height = image.size
        self._bg_photo = ctk.CTkImage(light_image=image, dark_image=image, size=(width, height))
        self.bg_label.configure(image=self._bg_photo)
        self.bg_label.lower()

    def _shimmer(self) -> None:
        if not self.winfo_exists():
            return
        self._phase = (self._phase + 0.045) % (math.tau)
        color = shimmer_color(self._phase)
        for card in self._cards:
            try:
                card.configure(border_color=color)
            except Exception:
                return
        self._show_background()
        self.after(120, self._shimmer)

    def _card(self, radius: int = 26) -> ctk.CTkFrame:
        card = ctk.CTkFrame(
            self,
            fg_color=GLASS,
            corner_radius=radius,
            border_width=2,
            border_color=MINT,
        )
        self._cards.append(card)
        return card

    def _build(self) -> None:
        self.bg_label = ctk.CTkLabel(self, text="")
        self.bg_label.place(x=0, y=0, relwidth=1, relheight=1)

        self.header = self._card(22)
        self.header.place(relx=0.02, rely=0.025, relwidth=0.96, relheight=0.1)
        if self.logo_image is not None:
            ctk.CTkLabel(self.header, text="", image=self.logo_image).pack(side="left", padx=(18, 8), pady=12)
        titles = ctk.CTkFrame(self.header, fg_color="transparent")
        titles.pack(side="left", fill="y", pady=10)
        ctk.CTkLabel(titles, text="Калькулятор клининга", text_color=TEXT, font=self._font(26, "bold"), anchor="w").pack(
            anchor="w"
        )
        ctk.CTkLabel(
            titles,
            text="Себестоимость, накладные и цена для клиента",
            text_color=MUTED,
            font=self._font(13),
            anchor="w",
        ).pack(anchor="w")
        self.refresh_button = ctk.CTkButton(
            self.header,
            text="Обновить данные",
            command=self.reload_data,
            width=180,
            height=40,
            corner_radius=20,
            fg_color="transparent",
            hover_color="#134e4a",
            border_width=2,
            border_color=MINT,
            text_color=MINT,
            font=self._font(14, "bold"),
        )
        self.refresh_button.pack(side="right", padx=18, pady=16)

        self.left = self._card()
        self.left.place(relx=0.02, rely=0.15, relwidth=0.46, relheight=0.82)
        self.right = self._card()
        self.right.place(relx=0.50, rely=0.15, relwidth=0.48, relheight=0.82)
        self._build_form()
        self._build_results()

    def _field_label(self, parent, text: str, row: int) -> None:
        ctk.CTkLabel(parent, text=text, text_color=MUTED, font=self._font(13), anchor="w").grid(
            row=row, column=0, sticky="w", pady=(10, 2)
        )

    def _build_form(self) -> None:
        form = ctk.CTkFrame(self.left, fg_color="transparent")
        form.pack(fill="both", expand=True, padx=20, pady=18)
        form.grid_columnconfigure(0, weight=1)

        self._field_label(form, "Площадь помещения (кв.м)", 0)
        self.area_entry = ctk.CTkEntry(
            form,
            textvariable=self.area_var,
            height=42,
            corner_radius=14,
            fg_color=FIELD,
            border_color=TEAL,
            text_color=TEXT,
            font=self._font(16),
            placeholder_text="например, 50",
        )
        self.area_entry.grid(row=1, column=0, sticky="ew")
        self.area_entry.bind("<Return>", lambda _event: self.calculate())

        self._field_label(form, "Тип уборки", 2)
        self.service_box = ctk.CTkComboBox(
            form,
            variable=self.service_var,
            values=[],
            state="readonly",
            height=42,
            corner_radius=14,
            fg_color=FIELD,
            border_color=TEAL,
            button_color=TEAL,
            button_hover_color=TEAL_DEEP,
            dropdown_fg_color=NAVY,
            dropdown_hover_color=TEAL_DEEP,
            dropdown_text_color=TEXT,
            text_color=TEXT,
            font=self._font(15),
            command=self._update_service_hint,
        )
        self.service_box.grid(row=3, column=0, sticky="ew")
        ctk.CTkLabel(form, textvariable=self.service_hint, text_color=MINT_DIM, font=self._font(12), anchor="w").grid(
            row=4, column=0, sticky="w", pady=(4, 0)
        )

        self._field_label(form, "Локация", 5)
        self.location_box = ctk.CTkComboBox(
            form,
            variable=self.location_var,
            values=["Центр", "Область"],
            state="readonly",
            height=42,
            corner_radius=14,
            fg_color=FIELD,
            border_color=TEAL,
            button_color=TEAL,
            button_hover_color=TEAL_DEEP,
            dropdown_fg_color=NAVY,
            dropdown_hover_color=TEAL_DEEP,
            dropdown_text_color=TEXT,
            text_color=TEXT,
            font=self._font(15),
            command=lambda _value: self._update_location_hint(),
        )
        self.location_box.grid(row=6, column=0, sticky="ew")
        ctk.CTkLabel(form, textvariable=self.location_hint, text_color=MINT_DIM, font=self._font(12), anchor="w").grid(
            row=7, column=0, sticky="w", pady=(4, 2)
        )

        flags = ctk.CTkFrame(form, fg_color="transparent")
        flags.grid(row=8, column=0, sticky="w", pady=(8, 4))
        self._check(flags, "Выходной день", self.weekend_var).pack(side="left", padx=(0, 18))
        self._check(flags, "Праздничный день", self.holiday_var).pack(side="left")

        ctk.CTkLabel(form, text="Дополнительные услуги", text_color=TEXT, font=self._font(16, "bold"), anchor="w").grid(
            row=9, column=0, sticky="w", pady=(12, 6)
        )
        self.extras = ctk.CTkScrollableFrame(
            form,
            fg_color=GLASS_SOFT,
            corner_radius=16,
            border_width=1,
            border_color="#245c56",
            scrollbar_button_color=TEAL,
            scrollbar_button_hover_color=MINT,
            height=230,
        )
        self.extras.grid(row=10, column=0, sticky="nsew")
        form.grid_rowconfigure(10, weight=1)

        self.calc_button = ctk.CTkButton(
            form,
            text="РАССЧИТАТЬ СТОИМОСТЬ",
            command=self.calculate,
            height=56,
            corner_radius=28,
            fg_color=TEAL,
            hover_color=TEAL_DEEP,
            border_width=2,
            border_color=MINT,
            text_color="#f0fdfa",
            font=self._font(18, "bold"),
        )
        self.calc_button.grid(row=11, column=0, sticky="ew", pady=(14, 0))
        ctk.CTkLabel(form, textvariable=self.status_var, text_color=MUTED, font=self._font(11), anchor="w", wraplength=460).grid(
            row=12, column=0, sticky="w", pady=(8, 0)
        )

    def _check(self, parent, text: str, variable: tk.BooleanVar) -> ctk.CTkCheckBox:
        return ctk.CTkCheckBox(
            parent,
            text=text,
            variable=variable,
            fg_color=TEAL,
            hover_color="#14b8a6",
            border_color=MINT,
            checkmark_color="#042f2e",
            text_color=TEXT,
            font=self._font(14),
        )

    def _build_results(self) -> None:
        box = ctk.CTkFrame(self.right, fg_color="transparent")
        box.pack(fill="both", expand=True, padx=22, pady=20)
        box.grid_columnconfigure(0, weight=1)
        box.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(box, text="Результат", text_color=TEXT, font=self._font(22, "bold"), anchor="w").grid(
            row=0, column=0, sticky="w"
        )
        self.placeholder = ctk.CTkLabel(
            box,
            text="Результаты появятся после расчёта.\nУкажите площадь, тип уборки и нажмите кнопку.",
            text_color=MUTED,
            font=self._font(16),
            justify="left",
            anchor="nw",
        )
        self.placeholder.grid(row=1, column=0, sticky="nw", pady=(18, 0))

        self.results_box = ctk.CTkFrame(box, fg_color="transparent")
        self.results_box.grid(row=1, column=0, sticky="nsew", pady=(12, 0))
        self.results_box.grid_columnconfigure(0, weight=1)
        self.results_box.grid_remove()
        self._metric(self.results_box, 0, "💰  Стоимость работ и материалов", self.base_var, self.base_detail, False)
        self._metric(self.results_box, 1, "📦  Накладные издержки", self.overhead_var, self.overhead_detail, False)
        self._metric(self.results_box, 2, "📈  Прибыль компании", self.markup_var, self.markup_detail, False)
        self.total_value = self._metric(self.results_box, 3, "🎯  ИТОГО К ОПЛАТЕ", self.total_var, None, True)

    def _metric(self, parent, row: int, caption: str, value: tk.StringVar, detail: tk.StringVar | None, accent: bool):
        card = ctk.CTkFrame(
            parent,
            fg_color="#12443c" if accent else "#143246",
            corner_radius=18,
            border_width=2 if accent else 1,
            border_color=MINT if accent else "#3d8f86",
        )
        card.grid(row=row, column=0, sticky="ew", pady=6)
        ctk.CTkLabel(
            card,
            text=caption,
            text_color=MINT if accent else MUTED,
            font=self._font(13, "bold" if accent else "normal"),
            anchor="w",
        ).pack(anchor="w", padx=16, pady=(12, 0))
        value_label = ctk.CTkLabel(
            card,
            textvariable=value,
            text_color=MINT if accent else TEXT,
            font=self._font(34 if accent else 26, "bold"),
            anchor="w",
        )
        value_label.pack(anchor="w", padx=16, pady=(0, 4 if detail is None else 0))
        if detail is not None:
            ctk.CTkLabel(card, textvariable=detail, text_color=MUTED, font=self._font(12), anchor="w", wraplength=500).pack(
                anchor="w", padx=16, pady=(0, 12)
            )
        else:
            value_label.pack_configure(pady=(0, 14))
        return value_label

    def reload_data(self, announce: bool = True) -> None:
        previous_service = self.service_var.get()
        previous_extras = {row.key: (row.checked.get(), row.quantity.get()) for row in self.extra_rows}
        try:
            catalog = load_catalog()
        except DataError as exc:
            self.catalog = None
            self.status_var.set("Данные не загружены")
            messagebox.showerror("Ошибка данных", str(exc), parent=self)
            return
        except Exception as exc:
            self.catalog = None
            self.status_var.set("Данные не загружены")
            messagebox.showerror("Ошибка данных", _friendly_io_error(exc).args[0], parent=self)
            return

        self.catalog = catalog
        names = [service.type_name for service in catalog.services]
        self.service_box.configure(values=names)
        self.service_var.set(previous_service if previous_service in names else names[0])
        self._update_service_hint()
        self._update_location_hint()
        self._rebuild_extras(previous_extras)
        self._hide_results()
        extras_count = len(extra_services(catalog))
        loaded = catalog.loaded_at.strftime("%H:%M:%S")
        self.status_var.set(
            f"Данные обновлены в {loaded}: {len(catalog.materials)} материалов, "
            f"{len(catalog.services)} типов услуг, {extras_count} доп. услуг."
        )
        if announce:
            messagebox.showinfo("Данные обновлены", "Файл calculator_data.xlsx прочитан заново.", parent=self)

    def _rebuild_extras(self, previous: dict[str, tuple[bool, str]]) -> None:
        for child in list(self.extras.winfo_children()):
            child.destroy()
        self.extra_rows = []
        if self.catalog is None:
            return
        services = extra_services(self.catalog)
        if not services:
            ctk.CTkLabel(
                self.extras,
                text="В листе «Лист3» нет дополнительных услуг.",
                text_color=MUTED,
                font=self._font(13),
                wraplength=420,
                justify="left",
            ).pack(anchor="w", padx=8, pady=8)
            return
        for item in services:
            checked_value, quantity_value = previous.get(item.key, (False, "1"))
            checked = tk.BooleanVar(value=checked_value)
            quantity = tk.StringVar(value=quantity_value or "1")
            row = ctk.CTkFrame(self.extras, fg_color="transparent")
            row.pack(fill="x", padx=(4, 12), pady=3)
            row.grid_columnconfigure(0, weight=1)
            self._check(row, item.description, checked).grid(row=0, column=0, sticky="w")
            entry = ctk.CTkEntry(
                row,
                textvariable=quantity,
                width=58,
                height=32,
                corner_radius=10,
                fg_color=FIELD,
                border_color=TEAL,
                text_color=TEXT,
                font=self._font(13),
                justify="center",
            )
            entry.grid(row=0, column=1, padx=(8, 4))
            ctk.CTkLabel(row, text=item.unit or "шт.", text_color=MUTED, font=self._font(12)).grid(row=0, column=2)
            ctk.CTkLabel(row, text=money(item.value), text_color=MINT, font=self._font(13, "bold")).grid(
                row=0, column=3, padx=(8, 0)
            )
            self.extra_rows.append(ExtraRow(item.key, checked, quantity))

    def _selected_service(self) -> Service | None:
        if self.catalog is None:
            return None
        return next((item for item in self.catalog.services if item.type_name == self.service_var.get()), None)

    def _update_service_hint(self, _value: str | None = None) -> None:
        service = self._selected_service()
        if service is None:
            self.service_hint.set("")
            return
        self.service_hint.set(
            f"Ставка {money(service.hourly_rate)}/ч · {fmt_num(service.productivity, 1)} кв.м/ч · "
            f"коэффициент {fmt_num(service.complexity, 2)} · минимум {fmt_num(service.min_hours, 1)} ч"
        )

    def _update_location_hint(self) -> None:
        if self.catalog is None:
            self.location_hint.set("")
            return
        key = "Транспорт_Центр" if self.location_var.get() == "Центр" else "Транспорт_Область"
        item = self.catalog.by_key.get(key)
        if item is None:
            self.location_hint.set("")
            return
        self.location_hint.set(f"{item.description}: {money(item.value)}")

    def _hide_results(self) -> None:
        self.results_box.grid_remove()
        self.placeholder.grid()

    def _quantities(self) -> dict[str, float]:
        chosen: dict[str, float] = {}
        for row in self.extra_rows:
            if not row.checked.get():
                continue
            label = next(
                (item.description for item in extra_services(self.catalog) if item.key == row.key),
                row.key,
            ) if self.catalog else row.key
            chosen[row.key] = parse_number(row.quantity.get(), f"Количество: {label}")
        return chosen

    def calculate(self) -> None:
        if self.catalog is None:
            messagebox.showwarning("Нет данных", "Сначала загрузите calculator_data.xlsx.", parent=self)
            return
        try:
            area = parse_number(self.area_var.get(), "Площадь помещения (кв.м)")
            result = calculate(
                self.catalog,
                area,
                self.service_var.get(),
                self._quantities(),
                self.weekend_var.get(),
                self.holiday_var.get(),
                self.location_var.get(),
            )
        except DataError as exc:
            messagebox.showerror("Не удалось рассчитать", str(exc), parent=self)
            return
        except Exception as exc:
            messagebox.showerror("Не удалось рассчитать", str(exc), parent=self)
            return
        self._show_result(result)

    def _show_result(self, result: Result) -> None:
        self.base_var.set(money(result.base))
        self.overhead_var.set(money(result.overhead))
        self.markup_var.set(money(result.markup))
        self.total_var.set(money(result.total))
        size = 34 if len(self.total_var.get()) <= 16 else 26
        self.total_value.configure(font=self._font(size, "bold"))
        self.base_detail.set(
            f"Материалы {money(result.materials)} · Работа {money(result.labor)} "
            f"({fmt_num(result.hours, 2)} ч) · Доп. услуги {money(result.extras)}"
        )
        self.overhead_detail.set(
            f"Транспорт {money(result.transport)} · Амортизация {money(result.amortization)} · "
            f"Налог {money(result.taxes)}"
        )
        self.markup_detail.set(
            f"Маржа {money(result.margin)} · Выходной {money(result.weekend_fee)} · "
            f"Праздник {money(result.holiday_fee)}"
        )
        self.placeholder.grid_remove()
        self.results_box.grid()
        self.status_var.set("Расчёт выполнен по текущему файлу calculator_data.xlsx.")


def main() -> None:
    if sys.platform.startswith("win"):
        try:
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("local.cleaning.calculator")
        except Exception:
            pass
    app = CalculatorApp()
    app.mainloop()


if __name__ == "__main__":
    main()
